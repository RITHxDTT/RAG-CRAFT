from openpyxl import load_workbook
from app.rag.loaders.base import ExtractedPage
from app.core.exceptions import AppError
from app.core.config import get_settings


def load(path):
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        count = 0
        for sheet in workbook.worksheets:
            # Reject pathological sparse dimensions before iterating millions of empty cells.
            if sheet.max_row and sheet.max_column and sheet.max_row * sheet.max_column > 2_000_000:
                raise AppError(422, 'This worksheet is too large. Split it into smaller files.')
            for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                values = [str(value) if value is not None else '' for value in row]
                if not any(values):
                    continue
                text = f'Sheet: {sheet.title}; Row: {number}\n' + ' | '.join(values)
                count += len(text)
                if count > get_settings().max_extracted_chars:
                    raise AppError(422, 'The extracted workbook is too large. Split it into smaller files.')
                yield ExtractedPage(text, sheet_name=sheet.title, row_number=number)
    finally:
        workbook.close()
