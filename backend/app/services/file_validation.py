from dataclasses import dataclass
from hashlib import sha256
from io import BytesIO
from pathlib import PurePosixPath
from zipfile import ZipFile
from docx import Document as DocxDocument
from pypdf import PdfReader
from app.core.exceptions import AppError

MIME_TYPES = {
    "xlsx": {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "application/zip"},
    "pdf": {"application/pdf"},
    "docx": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document", "application/zip"},
    "txt": {"text/plain"},
    "md": {"text/plain", "text/markdown", "text/x-markdown"},
}
CANONICAL_MIME = {"xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "pdf": "application/pdf", "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "txt": "text/plain", "md": "text/markdown"}


@dataclass(frozen=True)
class ValidatedFile:
    name: str
    file_type: str
    mime_type: str
    checksum: str


def validate_file(filename: str | None, mime: str | None, content: bytes, max_bytes: int) -> ValidatedFile:
    if not filename:
        raise AppError(422, "A filename is required.")
    name = PurePosixPath(filename.replace("\\", "/")).name
    extension = PurePosixPath(name).suffix.lower().lstrip(".")
    if not name or len(name) > 255 or any(ord(char) < 32 or ord(char) == 127 for char in name):
        raise AppError(422, "Use a valid filename of at most 255 characters.")
    if extension not in MIME_TYPES:
        raise AppError(415, "Supported file types are PDF, DOCX, TXT, MD, and XLSX.")
    if not content:
        raise AppError(422, "The file is empty.")
    if len(content) > max_bytes:
        raise AppError(413, "The file exceeds the configured upload size limit.")
    declared_mime = (mime or "").split(";")[0].strip().lower()
    if declared_mime not in MIME_TYPES[extension] | {"", "application/octet-stream"}:
        raise AppError(415, "The file's content type does not match its extension.")
    try:
        if extension == "pdf":
            if not content.startswith(b"%PDF-"):
                raise ValueError("Not a PDF")
            reader = PdfReader(BytesIO(content), strict=True)
            if reader.is_encrypted:
                raise AppError(422, "Password-protected PDFs are not supported. Upload an unlocked PDF.")
            if not reader.pages:
                raise ValueError("No pages")
        elif extension in {"docx", "xlsx"}:
            with ZipFile(BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > 10000 or sum(entry.file_size for entry in entries) > max_bytes * 5:
                    raise AppError(422, "The document expands beyond the permitted document size.")
                if ("word/document.xml" if extension == "docx" else "xl/workbook.xml") not in archive.namelist():
                    raise ValueError("Missing document")
                if any(entry.flag_bits & 1 for entry in entries):
                    raise ValueError("Encrypted archive")
                if archive.testzip() is not None:
                    raise ValueError("Corrupt archive")
            if extension == "docx":
                DocxDocument(BytesIO(content))
            else:
                from openpyxl import load_workbook
                workbook = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
                workbook.close()
        else:
            text = content.decode("utf-8-sig")
            if not text.strip():
                raise AppError(422, "The document contains no text.")
            if any(ord(char) < 32 and char not in "\t\n\r\f" for char in text):
                raise ValueError("Binary content")
    except AppError:
        raise
    except Exception:
        raise AppError(422, "The file is corrupted or unreadable. TXT and MD files must use UTF-8 encoding.") from None
    return ValidatedFile(name, extension.upper(), CANONICAL_MIME[extension], sha256(content).hexdigest())
