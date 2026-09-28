from io import BytesIO
import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from app.core.exceptions import AppError
from app.rag.chunker import chunk_pages
from app.rag.cleaner import clean_text
from app.rag.loaders import extract
from app.rag.loaders.base import ExtractedPage


def make_pdf(texts):
    writer = PdfWriter()
    for text in texts:
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(f'BT /F1 12 Tf 50 700 Td ({text}) Tj ET'.encode())
        page[NameObject('/Contents')] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def test_pdf_preserves_pages(tmp_path):
    path = tmp_path / 'handbook.pdf'
    path.write_bytes(make_pdf(['Annual leave is 18 days.', 'The office opens at nine.']))
    pages = extract(path, 'PDF')
    assert [p.page_number for p in pages] == [1, 2]
    chunks = chunk_pages(pages)
    assert '18 days' in chunks[0].content and chunks[0].page_number == 1
    assert chunks[1].page_number == 2


def test_docx_includes_tables(tmp_path):
    document = Document()
    document.add_paragraph('Leave policy')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Annual leave'
    table.cell(0, 1).text = '18 days'
    path = tmp_path / 'policy.docx'
    document.save(path)
    chunks = chunk_pages(extract(path, 'DOCX'))
    assert '18 days' in chunks[0].content and chunks[0].page_number is None


@pytest.mark.parametrize('kind', ['TXT', 'MD'])
def test_unicode_text(tmp_path, kind):
    path = tmp_path / 'policy.txt'
    path.write_text('សួស្តី\n\nAnnual leave: 18 days.', encoding='utf-8-sig')
    assert 'សួស្តី' in chunk_pages(extract(path, kind))[0].content


def test_chunk_size_overlap_and_full_coverage():
    text = ''.join(chr(0x400 + i) for i in range(300))
    chunks = chunk_pages([ExtractedPage(text, 3)], size=100, overlap=20)
    assert all(len(c.content) <= 100 and c.page_number == 3 for c in chunks)
    assert chunks[0].content[-20:] == chunks[1].content[:20]
    assert chunks[0].content + ''.join(c.content[20:] for c in chunks[1:]) == text
    assert [c.chunk_index for c in chunks] == list(range(len(chunks)))


def test_cleaner_is_conservative():
    assert clean_text('  Leave:  18 days.\r\n\n\n Next\x00 ') == 'Leave: 18 days.\n\nNext'


def test_unreadable_and_scanned_documents_fail(tmp_path):
    path = tmp_path / 'empty.pdf'
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    with path.open('wb') as file:
        writer.write(file)
    with pytest.raises(AppError, match='OCR'):
        extract(path, 'PDF')
    with pytest.raises(AppError, match='extraction failed'):
        extract(tmp_path / 'missing.txt', 'TXT')
    with pytest.raises(ValueError):
        chunk_pages([ExtractedPage('abc')], size=100, overlap=100)
